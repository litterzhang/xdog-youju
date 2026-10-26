"""Bounded acquisition-time document parsing; call only inside the host OS sandbox.

This module never fetches links, resolves external relationships, evaluates formulas,
or runs macros. Limits complement (not replace) the host's CPU/memory/time limits.
Locators refer to immutable input material; whitespace is normalized, not inferred.
"""

from __future__ import annotations

import codecs
import contextlib
import io
import json
import math
import os
import re
import signal
import struct
import subprocess
import tempfile
import zipfile
from datetime import date, datetime, time
from pathlib import PurePosixPath
from urllib.parse import unquote, urljoin, urlsplit
from xml.etree.ElementTree import ParseError

from bs4 import BeautifulSoup, Comment, NavigableString, Tag
from defusedxml import ElementTree as SafeET
from defusedxml.common import DefusedXmlException

from .regions import observe_regions, region_kind

MAX_INPUT_BYTES = 12 * 1024 * 1024
MAX_DECOMPRESSED_BYTES = 40 * 1024 * 1024
MAX_ENTRY_BYTES = 16 * 1024 * 1024
MAX_ZIP_ENTRIES = 1024
MAX_COMPRESSION_RATIO = 200
MAX_PAGES = 200
MAX_SHEETS = 50
MAX_ROWS = 5000
MAX_COLUMNS = 100
MAX_CELLS = 20000
MAX_TABLES = 200
MAX_BLOCKS = 24000
MAX_LINKS = 5000
MAX_NODES = 150000
MAX_TEXT_CHARS = 1000000
MAX_BLOCK_CHARS = 16000
MAX_OUTPUT_BYTES = 8 * 1024 * 1024
MAX_PDF_STREAM_BYTES = 4 * 1024 * 1024
MAX_ERRORS = 64
ANTIWORD_TIMEOUT = 10

_SUFFIXES = {
    ".html": "html",
    ".htm": "html",
    ".xhtml": "html",
    ".pdf": "pdf",
    ".xlsx": "xlsx",
    ".xlsm": "xlsx",
    ".xls": "xls",
    ".docx": "docx",
    ".docm": "docx",
    ".doc": "doc",
    ".xml": "xml",
}
_MEDIA = {
    "text/html": "html",
    "application/xhtml+xml": "html",
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "application/vnd.ms-excel.sheet.macroenabled.12": "xlsx",
    "application/vnd.ms-excel": "xls",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.ms-word.document.macroenabled.12": "docx",
    "application/msword": "doc",
    "application/xml": "xml",
    "text/xml": "xml",
}
_ATTACH_SUFFIXES = set(_SUFFIXES) - {".html", ".htm", ".xhtml", ".xml"}
_ATTACH_SUFFIXES.update({".zip", ".rar", ".7z", ".csv", ".rtf", ".ppt", ".pptx", ".odt"})
_ATTACHMENT_LABEL = re.compile(
    r"附件|下载|岗位表|职位表|报名表|attachment|download|\.(?:pdf|docx?|xlsx?|zip|csv|rtf)\b", re.I
)
_FILE_LABEL = re.compile(r"\.(?:pdf|docx?|xlsx?|zip|csv|rtf)\b", re.I)
_NAVIGATION_NAME = re.compile(
    r"(?:^|[-_])(?:nav|menu|navigation|navbar|menubar)(?:$|[-_]|\d)", re.I
)
_SPACE = re.compile(r"\s+")
_XML_ERRORS = (ParseError, DefusedXmlException)
_LINK_KIND_PRIORITY = {"asset": 0, "form": 1, "page": 2, "attachment": 3}
_ASSET_LINK_RELS = {
    "stylesheet",
    "icon",
    "apple-touch-icon",
    "apple-touch-icon-precomposed",
    "mask-icon",
    "manifest",
    "modulepreload",
    "dns-prefetch",
    "preconnect",
}
_PAGE_LINK_RELS = {
    "alternate",
    "canonical",
    "next",
    "prev",
    "first",
    "last",
    "up",
    "index",
    "contents",
    "chapter",
    "section",
    "subsection",
    "start",
    "help",
    "author",
    "license",
    "search",
}


class _MaterialError(Exception):
    def __init__(self, code: str, message: str, *, truncated: bool = False):
        super().__init__(message)
        self.code, self.truncated = code, truncated


class _Limit(_MaterialError):
    def __init__(self, message: str):
        super().__init__("LIMIT_EXCEEDED", message, truncated=True)


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and math.isfinite(value) and value.is_integer():
        return str(int(value))
    return _SPACE.sub(" ", str(value)).strip()


class _Document:
    def __init__(self, material_key: str, format_: str):
        self.key = material_key
        self.data = {
            "format": format_,
            "status": "ok",
            "title": "",
            "blocks": [],
            "tables": [],
            "links": [],
            "errors": [],
            "truncated": False,
        }
        self.text_chars = 0
        self.cells = 0
        self._block_text: dict[str, str] = {}
        self._links: dict[str, dict] = {}

    def issue(self, code: str, message: str, *, truncated: bool = False):
        self.data["status"] = "partial"
        self.data["truncated"] |= truncated
        # A representative error per code, not unbounded per-cell diagnostics.
        errors = self.data["errors"]
        if not any(item["code"] == code for item in errors):
            if len(errors) < MAX_ERRORS:
                errors.append({"code": code, "message": message[:1200]})
            else:
                self.data["truncated"] = True
                errors[-1] = {"code": "ERROR_LIMIT", "message": "Further diagnostics omitted."}

    def title(self, value):
        value = _text(value)
        if len(value) > 2000:
            self.issue("TITLE_TRUNCATED", "The title was truncated.", truncated=True)
        self.data["title"] = value[:2000]

    def block(self, value, locator: str) -> str:
        if len(self.data["blocks"]) >= MAX_BLOCKS:
            raise _Limit("Material block limit exceeded.")
        value = _text(value)
        if len(value) > MAX_BLOCK_CHARS:
            self.issue("TEXT_TRUNCATED", "An individual text block was truncated.", truncated=True)
            value = value[:MAX_BLOCK_CHARS]
        if self.text_chars + len(value) > MAX_TEXT_CHARS:
            raise _Limit("Material text limit exceeded.")
        self.text_chars += len(value)
        key = f"{self.key}:b{len(self.data['blocks']) + 1}"
        self.data["blocks"].append({"id": key, "text": value, "locator": locator})
        self._block_text[key] = value
        return key

    def table(self, locator: str, evidence: list[list[str]]):
        if len(self.data["tables"]) >= MAX_TABLES:
            raise _Limit("Material table limit exceeded.")
        size = sum(map(len, evidence))
        if self.cells + size > MAX_CELLS:
            raise _Limit("Material normalized cell limit exceeded.")
        self.cells += size
        self.data["tables"].append(
            {
                "id": f"{self.key}:t{len(self.data['tables']) + 1}",
                "locator": locator,
                "rows": [[self._block_text[key] for key in row] for row in evidence],
                "cell_evidence": evidence,
            }
        )

    def link(
        self,
        raw: str,
        *,
        base: str,
        label: str = "",
        attachment: bool = False,
        infer_label: bool = True,
        kind: str = "page",
    ):
        if kind not in _LINK_KIND_PRIORITY:
            raise ValueError("Unknown discovery link kind")
        if not isinstance(raw, str) or not raw.strip():
            return
        raw = raw.strip()
        if len(raw) > 8192:
            self.issue("LINK_TRUNCATED", "An excessively long URL was omitted.", truncated=True)
            return
        try:
            url = urljoin(base, raw)
            parts = urlsplit(url)
            if parts.scheme.lower() not in {"http", "https"} or not parts.hostname:
                return
            _ = parts.port  # Reject invalid syntactic ports, not private/scope targets.
        except (ValueError, UnicodeError):
            self.issue("INVALID_LINK", "An invalid URL was omitted.")
            return
        label = _text(label)
        if kind in {"page", "attachment"}:
            attachment |= kind == "attachment"
            attachment |= PurePosixPath(unquote(parts.path).lower()).suffix in _ATTACH_SUFFIXES
            # Filename/download hints apply only to document relationships. An
            # image named .pdf or a form action ending in .docx is not a download.
            attachment |= bool((_ATTACHMENT_LABEL if infer_label else _FILE_LABEL).search(label))
            kind = "attachment" if attachment else "page"
        else:
            attachment = False
        if len(label) > 2000:
            self.issue("LINK_LABEL_TRUNCATED", "A link label was truncated.", truncated=True)
        value = {
            "url": url,
            "raw_url": raw,
            "label": label[:2000],
            "is_attachment": attachment,
            "kind": kind,
        }
        if url in self._links:
            existing = self._links[url]
            if _LINK_KIND_PRIORITY[kind] > _LINK_KIND_PRIORITY[existing["kind"]]:
                existing.update(value)
            elif kind == existing["kind"] and not existing["label"]:
                existing["label"] = value["label"]
            return
        if len(self._links) >= MAX_LINKS:
            self.issue("LINK_LIMIT", "Additional links were omitted.", truncated=True)
            return
        self._links[url] = value
        self.data["links"].append(value)

    def finish(self) -> dict:
        # Keep the contract/evidence integrity instead of cutting serialized JSON.
        if len(json.dumps(self.data, ensure_ascii=True).encode("utf-8")) > MAX_OUTPUT_BYTES:
            self.data["blocks"].clear()
            self.data["tables"].clear()
            self.data["links"].clear()
            self.data.pop("structure", None)  # No dangling region IDs after bounded-output failure.
            self.issue(
                "OUTPUT_LIMIT", "Material output exceeded the JSON byte limit.", truncated=True
            )
            self.data["status"] = "failed"
        return self.data


def _suffix(url: str) -> str:
    try:
        return PurePosixPath(unquote(urlsplit(url).path).lower()).suffix
    except (ValueError, UnicodeError):
        return ""


def _format(body: bytes, media_type: str, url: str) -> str:
    head = body[:1024].lstrip(b"\xef\xbb\xbf \r\n\t")
    if head.startswith(b"%PDF-"):
        return "pdf"
    if head.lower().startswith((b"<!doctype html", b"<html", b"<head", b"<body")):
        return "html"
    if head.startswith(b"<?xml"):
        if (
            b"<html" in head.lower()
            or media_type.split(";", 1)[0].strip().lower() == "application/xhtml+xml"
        ):
            return "html"
        return "xml"
    if re.match(rb"<(?:\w+:)?(?:urlset|sitemapindex)\b", head):
        return "xml"
    hint = _MEDIA.get(media_type.split(";", 1)[0].strip().lower(), _SUFFIXES.get(_suffix(url)))
    if body.startswith(b"PK\x03\x04"):
        return hint if hint in {"xlsx", "docx"} else "zip"
    if body.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return hint if hint in {"doc", "xls"} else "unknown"
    return hint or "unknown"


def parse_material(body: bytes, *, media_type: str, url: str, material_key: str) -> dict:
    """Return the exact material Document contract; inputs are bytes and host identities.

    Format/resource failures are data. Caller programming errors (invalid argument
    types or broken library APIs) are not disguised as successful empty documents.
    The caller must enforce its own sandbox deadline and translate worker failures.
    """
    if not isinstance(body, bytes):
        raise TypeError("body must be bytes")
    if not isinstance(material_key, str) or not material_key or len(material_key) > 256:
        raise ValueError("material_key must be a nonempty, bounded host version identity")
    result = _Document(material_key, _format(body, media_type, url))
    try:
        if len(body) > MAX_INPUT_BYTES:
            raise _Limit("Material input byte limit exceeded.")
        if not body:
            raise _MaterialError("EMPTY_INPUT", "The material has no bytes.")
        format_ = result.data["format"]
        if format_ in {"zip", "xlsx", "docx"}:
            with _checked_zip(body, result) as package:
                names = set(package.namelist())
                if "word/document.xml" in names:
                    format_ = "docx"
                elif "xl/workbook.xml" in names:
                    format_ = "xlsx"
                else:
                    format_ = "unknown"
                result.data["format"] = format_
                if format_ == "xlsx":
                    _xlsx(body, result)
                elif format_ == "docx":
                    _docx(body, result)
        elif format_ == "html":
            _html(body, url, result, media_type)
        elif format_ == "xml":
            _sitemap(body, url, result)
        elif format_ == "pdf":
            _pdf(body, result)
        elif format_ == "xls":
            _xls(body, result)
        elif format_ == "doc":
            _doc(body, result)
        if format_ == "unknown":
            result.issue("UNSUPPORTED_FORMAT", "This material format is not supported.")
            result.data["status"] = "unsupported"
    except _MaterialError as exc:
        result.issue(exc.code, str(exc), truncated=exc.truncated)
        if not result.data["blocks"] and not result.data["links"]:
            result.data["status"] = "failed"
    return result.finish()


@contextlib.contextmanager
def _checked_zip(body: bytes, result: _Document):
    """Validate every member before any OOXML library processes compressed input."""
    try:
        with zipfile.ZipFile(io.BytesIO(body)) as package:
            infos = package.infolist()
            if len(infos) > MAX_ZIP_ENTRIES:
                raise _Limit("Too many ZIP entries.")
            if len({entry.filename for entry in infos}) != len(infos):
                raise _MaterialError("INVALID_ZIP", "Duplicate ZIP member names are ambiguous.")
            if sum(entry.file_size for entry in infos) > MAX_DECOMPRESSED_BYTES:
                raise _Limit("ZIP expanded size exceeds the limit.")
            expanded = 0
            for entry in infos:
                if entry.flag_bits & 1:
                    raise _MaterialError(
                        "ENCRYPTED_DOCUMENT", "Encrypted ZIP entries are unsupported."
                    )
                if entry.file_size > MAX_ENTRY_BYTES:
                    raise _Limit("ZIP member expanded size exceeds the limit.")
                if entry.file_size > max(1, entry.compress_size) * MAX_COMPRESSION_RATIO:
                    raise _Limit("ZIP compression ratio exceeds the limit.")
                name = entry.filename
                if ".." in PurePosixPath(name).parts or name.startswith(("/", "\\")):
                    raise _MaterialError("INVALID_ZIP", "Unsafe ZIP member name.")
                if entry.compress_type not in {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED}:
                    raise _MaterialError("UNSUPPORTED_ZIP", "Unsupported ZIP compression method.")
                with package.open(entry) as stream:
                    data = stream.read(MAX_ENTRY_BYTES + 1)
                expanded += len(data)
                if len(data) > MAX_ENTRY_BYTES or expanded > MAX_DECOMPRESSED_BYTES:
                    raise _Limit("Actual ZIP expansion exceeds the limit.")
                if name.lower().endswith((".xml", ".rels")):
                    root = _safe_xml(data)
                    if sum(1 for _ in root.iter()) > MAX_NODES:
                        raise _Limit("OOXML element count exceeds the limit.")
                    for element in root.iter():
                        local_name = element.tag.rsplit("}", 1)[-1]
                        if local_name == "gridSpan":
                            value = next(
                                (
                                    v
                                    for k, v in element.attrib.items()
                                    if k.rsplit("}", 1)[-1] == "val"
                                ),
                                "1",
                            )
                            try:
                                span = int(value)
                            except ValueError as exc:
                                raise _MaterialError(
                                    "INVALID_DOCX", "Invalid Word merged-cell span."
                                ) from exc
                            if not 1 <= span <= MAX_COLUMNS:
                                raise _Limit("Word merged-cell span exceeds the limit.")
                        elif local_name == "mergeCell":
                            from openpyxl.utils.cell import range_boundaries

                            try:
                                left, top, right, bottom = range_boundaries(
                                    element.attrib.get("ref", "")
                                )
                                if None in (left, top, right, bottom):
                                    raise ValueError
                            except ValueError as exc:
                                raise _MaterialError(
                                    "INVALID_XLSX", "Invalid Excel merged-cell range."
                                ) from exc
                            if (right - left + 1) * (bottom - top + 1) > MAX_CELLS:
                                raise _Limit("Excel merged-cell expansion exceeds the limit.")
                    if name.lower().endswith(".rels"):
                        for rel in root:
                            if rel.attrib.get("TargetMode", "").lower() == "external":
                                result.issue(
                                    "EXTERNAL_RELATIONSHIP_IGNORED",
                                    "External relationships are not fetched.",
                                )
                if "vbaproject" in name.lower() or "/embeddings/" in name.lower():
                    result.issue(
                        "ACTIVE_CONTENT_IGNORED",
                        "Macros and embedded objects are not executed or parsed.",
                    )
            yield package
    except (zipfile.BadZipFile, EOFError) as exc:
        raise _MaterialError("INVALID_ZIP", "The Office ZIP container is corrupt.") from exc


def _safe_xml(body: bytes):
    try:
        return SafeET.fromstring(body, forbid_dtd=True, forbid_entities=True, forbid_external=True)
    except _XML_ERRORS as exc:
        raise _MaterialError(
            "UNSAFE_OR_INVALID_XML", "XML is malformed or contains forbidden declarations."
        ) from exc


def _locator(tag: Tag) -> str:
    parts = []
    while isinstance(tag, Tag) and tag.name != "[document]":
        if len(parts) >= 100:
            raise _Limit("HTML nesting exceeds the locator depth limit.")
        ordinal = (
            sum(
                1
                for node in tag.previous_siblings
                if isinstance(node, Tag) and node.name == tag.name
            )
            + 1
        )
        parts.append(f"{tag.name}:nth-of-type({ordinal})")
        tag = tag.parent
    return " > ".join(reversed(parts)) or "html:document"


def _in_navigation(tag: Tag) -> bool:
    """Recognize structural navigation, not source-specific URL or business rules."""
    for ancestor in (tag, *tag.parents):
        if ancestor.name in {"nav", "header", "footer"} or ancestor.get("role") in {
            "navigation",
            "menu",
            "menubar",
        }:
            return True
        names = [ancestor.get("id", ""), *ancestor.get("class", [])]
        if any(_NAVIGATION_NAME.search(name) for name in names):
            return True
    return False


def _layout_table(table: Tag, rows: list[Tag]) -> bool:
    """Flatten clear layout wrappers while preserving ambiguous/data tables.

    Nested content is not a second set of parent-table cells. A real table with
    its own header/caption/role remains tabular, even if one cell embeds a table.
    All text outside retained data tables is emitted separately at its original
    DOM location, so flattening layout does not discard announcement prose.
    """
    if table.get("role") in {"presentation", "none"}:
        return True
    cells = [row.find_all(["td", "th"], recursive=False) for row in rows]
    if (
        table.get("role") in {"table", "grid", "treegrid"}
        or table.find("caption", recursive=False)
        or any(
            cell.name == "th" or cell.has_attr("headers") or cell.has_attr("scope")
            for row in cells
            for cell in row
        )
    ):
        return False
    first = next((row for row in cells if any(cell.get_text(strip=True) for cell in row)), [])
    if not first:
        return True  # Empty decorative cells contain no literal tabular values.
    if table.find("table") and all(
        not cell.get_text(strip=True) or cell.find("table") for cell in first
    ):
        return True
    if _in_navigation(table):
        return True
    texts = [
        node
        for node in table.descendants
        if isinstance(node, NavigableString) and not isinstance(node, Comment) and _text(node)
    ]
    return bool(texts) and all(_in_navigation(node.parent) for node in texts)


def _html_link_kind(tag: Tag, attribute: str) -> str:
    """Classify the DOM relationship, never URL words or recruitment semantics.

    ``asset`` also covers non-navigable document support references such as a
    base URL, preconnect, or an unrecognized link relation. Inventory retains
    them, but the host must not enqueue them as pages or submit forms.
    """
    if attribute in {"action", "formaction"}:
        return "form"
    if attribute in {"poster", "srcset"}:
        return "asset"
    if attribute in {"cite", "longdesc"}:
        return "page"
    if attribute == "href":
        if tag.name in {"a", "area"}:
            return "page"
        if tag.name == "link":
            relations = {rel.lower() for rel in tag.get("rel", [])}
            if relations & _ASSET_LINK_RELS:
                return "asset"
            if relations & {"preload", "prefetch"}:
                return "page" if tag.get("as", "").lower() == "document" else "asset"
            return "page" if relations & _PAGE_LINK_RELS else "asset"
        return "asset"
    if (attribute == "src" and tag.name in {"iframe", "frame", "embed"}) or (
        attribute == "data" and tag.name == "object"
    ):
        media_type = tag.get("type", "").split(";", 1)[0].strip().lower()
        if media_type.startswith(("image/", "audio/", "video/")):
            return "asset"
        return "page"
    return "asset"


def _html(body: bytes, url: str, result: _Document, media_type: str):
    charset = re.search(r"charset\s*=\s*[\"']?([^\s;\"']+)", media_type, re.I)
    encoding = charset.group(1) if charset else None
    if encoding:
        try:
            codecs.lookup(encoding)
        except LookupError:
            result.issue("UNKNOWN_CHARSET", "Unknown HTTP charset; encoding detection was used.")
            encoding = None
    if encoding is None:
        try:
            body.decode("utf-8")
            encoding = "utf-8"
        except UnicodeDecodeError:
            pass  # Allow the parser to inspect declared legacy encodings.
    soup = BeautifulSoup(body, "html.parser", from_encoding=encoding)
    if soup.contains_replacement_characters:
        result.issue(
            "TEXT_DECODE_REPLACEMENT", "Undecodable bytes were replaced while reading HTML."
        )
    if sum(1 for _ in soup.descendants) > MAX_NODES:
        raise _Limit("HTML node count exceeds the limit.")
    result.title(soup.title.get_text(" ", strip=True) if soup.title else "")
    # Preserve non-body date observations separately; routing may approximate age.
    # These values are not authenticated publication dates or position deadlines.
    date_names = {
        "article:published_time",
        "article:modified_time",
        "og:updated_time",
        "pubdate",
        "publishdate",
        "publicationdate",
        "date",
        "datepublished",
        "datemodified",
        "dcterms.created",
        "dcterms.modified",
        "dc.date",
        "last-modified",
        "lastmodified",
        "updatetime",
        "createdate",
    }
    hints = []
    for tag in soup.find_all(["meta", "time"]):
        if len(hints) >= 64:
            break
        name = str(
            tag.get("property")
            or tag.get("name")
            or tag.get("itemprop")
            or tag.get("http-equiv")
            or ""
        ).lower()
        if tag.name == "meta" and name in date_names:
            value, kind = tag.get("content"), "html_meta"
        elif tag.name == "time" and not _in_navigation(tag) and region_kind(tag) == "content":
            value, kind = tag.get("datetime") or tag.get_text(" ", strip=True), "html_time"
        else:
            continue
        if isinstance(value, str) and value.strip() and len(value) <= 256:
            hints.append({"kind": kind, "value": value.strip(), "locator": _locator(tag)})
    if hints:
        result.data["date_hints"] = hints

    base = url
    base_tag = soup.find("base", href=True)
    if base_tag:
        try:
            candidate = urljoin(url, base_tag["href"])
            if urlsplit(candidate).scheme in {"http", "https"}:
                base = candidate
            else:
                result.issue("INVALID_BASE_URL", "Non-HTTP base URL ignored.")
        except (ValueError, UnicodeError):
            result.issue("INVALID_BASE_URL", "Malformed base URL ignored.")
    for tag in soup.find_all(True):
        label = (
            tag.get_text(" ", strip=True) if tag.name in {"a", "area", "button", "object"} else ""
        )
        label = label or tag.get("title", "") or tag.get("aria-label", "") or tag.get("alt", "")
        type_hint = tag.get("type", "").split(";", 1)[0].lower()
        attachment = (
            "download" in tag.attrs
            or type_hint
            in {"application/octet-stream", "application/zip", "application/rtf", "text/csv"}
            or _MEDIA.get(type_hint)
            in {
                "pdf",
                "docx",
                "doc",
                "xls",
                "xlsx",
            }
        )
        infer_label = not _in_navigation(tag)
        for attribute in ("href", "src", "action", "formaction", "poster", "cite", "longdesc"):
            if tag.has_attr(attribute):
                result.link(
                    tag[attribute],
                    base=base,
                    label=label,
                    attachment=attachment,
                    infer_label=infer_label,
                    kind=_html_link_kind(tag, attribute),
                )
                try:
                    linked = result._links.get(urljoin(base, tag[attribute]))
                    if linked is not None:
                        observed_kind = region_kind(tag)
                        # A content occurrence wins over an auxiliary duplicate.
                        if "region_kind" not in linked or observed_kind == "content":
                            linked["region_kind"] = observed_kind
                except (ValueError, TypeError):
                    pass  # Invalid URLs already have bounded parser diagnostics.
        if tag.name == "object" and tag.has_attr("data"):
            result.link(
                tag["data"],
                base=base,
                label=label,
                attachment=attachment,
                infer_label=infer_label,
                kind=_html_link_kind(tag, "data"),
            )
        if tag.has_attr("srcset"):
            for item in tag["srcset"].split(","):
                tokens = item.strip().split()
                if tokens:
                    result.link(
                        tokens[0], base=base, label=label, infer_label=infer_label, kind="asset"
                    )
    main = (
        soup.find("main")
        or soup.find("article")
        or soup.find(id="oInfoContent")
        or soup.body
        or soup
    )
    tables = []
    for table in main.find_all("table"):
        rows = [row for row in table.find_all("tr") if row.find_parent("table") is table]
        if not _layout_table(table, rows):
            tables.append((table, rows))
    data_table_ids = {id(table) for table, _ in tables}
    # Group literal text by its original parent block without pruning/mutating the DOM.
    groups: list[tuple[Tag, list[str]]] = []
    skip = {"script", "style", "template", "noscript", "head"}
    block_tags = {
        "p",
        "li",
        "dt",
        "dd",
        "pre",
        "blockquote",
        "div",
        "section",
        "article",
        "main",
        "body",
        "a",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "td",
        "th",
    }
    for node in main.descendants:
        if not isinstance(node, NavigableString) or isinstance(node, Comment) or not _text(node):
            continue
        ancestors = []
        for ancestor in node.parents:
            ancestors.append(ancestor)
            if ancestor is main:
                break
        if any(tag.name in skip or id(tag) in data_table_ids for tag in ancestors):
            continue
        parent = next((tag for tag in ancestors if tag.name in block_tags), node.parent)
        if groups and groups[-1][0] is parent:
            groups[-1][1].append(str(node))
        else:
            groups.append((parent, [str(node)]))
    for tag, parts in groups:
        result.block(" ".join(parts), _locator(tag))
        result.data["blocks"][-1]["region_kind"] = region_kind(tag)
    for table, rows in tables:
        if len(rows) > MAX_ROWS:
            result.issue("ROW_LIMIT", "HTML table rows were truncated.", truncated=True)
            rows = rows[:MAX_ROWS]
        grid: dict[tuple[int, int], str] = {}
        row_locations = [_locator(row) for row in rows]
        for ri, row in enumerate(rows):
            column = 0
            for cell in row.find_all(["td", "th"], recursive=False):
                while (ri, column) in grid:
                    column += 1
                rowspan = _span(cell.get("rowspan", 1), result)
                colspan = _span(cell.get("colspan", 1), result)
                if rowspan == 0:
                    rowspan = sum(1 for following in rows[ri:] if following.parent is row.parent)
                if colspan == 0:
                    colspan = 1
                    result.issue("INVALID_TABLE_SPAN", "Zero colspan was treated as one.")
                if column + colspan > MAX_COLUMNS or rowspan > MAX_ROWS:
                    raise _Limit("HTML table span exceeds the limit.")
                if ri + rowspan > len(rows):
                    result.issue("INVALID_TABLE_SPAN", "Rowspan exceeds available source rows.")
                block = result.block(cell.get_text(" ", strip=True), _locator(cell))
                for rr in range(ri, min(ri + rowspan, len(rows))):
                    for cc in range(column, column + colspan):
                        if (rr, cc) in grid:
                            result.issue(
                                "OVERLAPPING_TABLE_SPAN",
                                "Overlapping cells make this table ambiguous.",
                            )
                        else:
                            grid[rr, cc] = block
                        if len(grid) + result.cells > MAX_CELLS:
                            raise _Limit("HTML normalized table cell limit exceeded.")
                column += colspan
        width = max((col for _, col in grid), default=-1) + 1
        if width and rows:
            matrix = []
            occupied_rows = {ri for ri, _ in grid}
            for ri, row_loc in enumerate(row_locations):
                if ri not in occupied_rows:
                    # An empty <tr> used as visual spacing has no source cells.
                    # Keep rows occupied by rowspan, but do not invent gap cells.
                    continue
                values = []
                for ci in range(width):
                    key = grid.get((ri, ci))
                    if key is None:
                        # Empty structural slots reference the existing source row.
                        result.issue(
                            "SPARSE_TABLE",
                            "Missing HTML cells were represented as empty source-row slots.",
                        )
                        key = result.block("", row_loc)
                    values.append(key)
                matrix.append(values)
            result.table(_locator(table), matrix)
    result.data["structure"] = observe_regions(soup, result.data, _locator)


def _span(value, result: _Document) -> int:
    try:
        number = int(value)
        if number < 0:
            raise ValueError
        return number
    except (TypeError, ValueError):
        result.issue("INVALID_TABLE_SPAN", "An invalid merged-cell span was treated as one.")
        return 1


def _sitemap(body: bytes, url: str, result: _Document):
    root = _safe_xml(body)

    def local(name):
        return name.rsplit("}", 1)[-1]  # Namespace prefixes are not identities.

    if local(root.tag) not in {"urlset", "sitemapindex"}:
        result.issue("UNSUPPORTED_XML", "XML is not a URL sitemap or sitemap index.")
        result.data["status"] = "unsupported"
        return
    for index, element in enumerate(root.iter()):
        if index >= MAX_NODES:
            raise _Limit("Sitemap XML element limit exceeded.")
        if local(element.tag) == "loc" and element.text:
            result.link(element.text.strip(), base=url, label="sitemap entry")


@contextlib.contextmanager
def _pdf_limits():
    # Public pypdf configuration limits every decoding stage; the wrapper tracks
    # cumulative decoded content, including fonts and XObject Forms. Each parser
    # call is process-isolated by the host, not run concurrently in threads.
    from pypdf import apply_configuration, filters

    decode = filters.decode_stream_data
    expanded = 0

    def bounded_decode(stream):
        nonlocal expanded
        value = decode(stream)
        expanded += len(value)
        if len(value) > MAX_PDF_STREAM_BYTES or expanded > MAX_DECOMPRESSED_BYTES:
            raise _Limit("PDF cumulative stream expansion exceeds the limit.")
        return value

    try:
        filters.decode_stream_data = bounded_decode
        with apply_configuration(
            maximum_declared_stream_length=MAX_PDF_STREAM_BYTES,
            array_based_stream_maximum_output_length=MAX_PDF_STREAM_BYTES,
            zlib_maximum_output_length=MAX_PDF_STREAM_BYTES,
            zlib_maximum_recovery_input_length=MAX_PDF_STREAM_BYTES,
            lzw_maximum_output_length=MAX_PDF_STREAM_BYTES,
            run_length_maximum_output_length=MAX_PDF_STREAM_BYTES,
            brotli_maximum_output_length=MAX_PDF_STREAM_BYTES,
            jbig2_maximum_output_length=MAX_PDF_STREAM_BYTES,
            image_maximum_buffer_size=MAX_PDF_STREAM_BYTES,
            flate_maximum_row_length=MAX_PDF_STREAM_BYTES,
            stream_decoding_work_maximum_length=MAX_DECOMPRESSED_BYTES,
            stream_filters_maximum_length=8,
            page_tree_maximum_entries=10000,
            xform_maximum_invocations_per_extraction=1000,
            jbig2dec_binary=None,
        ):
            yield
    finally:
        filters.decode_stream_data = decode


def _pdf(body: bytes, result: _Document):
    from pypdf import PdfReader
    from pypdf.errors import DependencyError, LimitReachedError, PyPdfError

    try:
        with _pdf_limits():
            reader = PdfReader(io.BytesIO(body), strict=True)
            if reader.is_encrypted:
                raise _MaterialError(
                    "ENCRYPTED_PDF", "Encrypted PDFs require review; passwords are not attempted."
                )
            if reader.metadata and reader.metadata.title:
                result.title(reader.metadata.title)
            count = len(reader.pages)
            if count > MAX_PAGES:
                result.issue("PAGE_LIMIT", "PDF pages were truncated.", truncated=True)
            expanded = 0
            for index in range(min(count, MAX_PAGES)):
                page = reader.pages[index]
                contents = page.get_contents()
                if contents is not None:
                    size = len(contents.get_data())
                    expanded += size
                    if size > MAX_PDF_STREAM_BYTES or expanded > MAX_DECOMPRESSED_BYTES:
                        raise _Limit("PDF decoded content stream limit exceeded.")
                text = page.extract_text()
                if len(text) > MAX_TEXT_CHARS:
                    raise _Limit("A PDF page exceeds the text limit.")
                if not text.strip():
                    result.issue(
                        "PDF_NO_TEXT",
                        f"PDF page {index + 1} has no extractable text; "
                        "it may be scanned or blank. OCR is not performed.",
                    )
                    continue
                # Real page/line references, never guessed table cells or page numbers.
                for line_number, line in enumerate(text.splitlines(), 1):
                    if line.strip():
                        result.block(line, f"pdf:page={index + 1};text-line={line_number}")
            if not count:
                result.issue("PDF_NO_PAGES", "The PDF has no pages.")
    except LimitReachedError as exc:
        raise _Limit("PDF decoder expansion limit exceeded.") from exc
    except (PyPdfError, DependencyError, NotImplementedError) as exc:
        raise _MaterialError(
            "INVALID_OR_UNSUPPORTED_PDF", "PDF parsing failed or requires an unsupported codec."
        ) from exc


def _merge_map(ranges, rows: int, columns: int, result: _Document):
    merged = {}
    for first_row, end_row, first_column, end_column in ranges:
        if end_row - first_row > MAX_ROWS or end_column - first_column > MAX_COLUMNS:
            result.issue("MERGE_LIMIT", "An oversized merged range was clipped.", truncated=True)
        for row in range(max(0, first_row), min(rows, end_row)):
            for column in range(max(0, first_column), min(columns, end_column)):
                if (row, column) in merged:
                    result.issue(
                        "OVERLAPPING_TABLE_SPAN",
                        "Overlapping spreadsheet merge ranges are ambiguous.",
                    )
                merged[row, column] = (first_row, first_column)
                if len(merged) > MAX_CELLS:
                    raise _Limit("Spreadsheet merged-cell count exceeded.")
    return merged


def _sheet_size(rows: int, columns: int, result: _Document) -> tuple[int, int]:
    if rows > MAX_ROWS:
        result.issue("ROW_LIMIT", "Spreadsheet rows were truncated.", truncated=True)
        rows = MAX_ROWS
    if columns > MAX_COLUMNS:
        result.issue("COLUMN_LIMIT", "Spreadsheet columns were truncated.", truncated=True)
        columns = MAX_COLUMNS
    remaining = MAX_CELLS - result.cells
    if rows * columns > remaining:
        result.issue(
            "CELL_LIMIT",
            "Spreadsheet rows were limited by the normalized cell budget.",
            truncated=True,
        )
        rows = remaining // max(columns, 1)
    return rows, columns


def _xlsx_data_bounds(sheet) -> tuple[int, int]:
    """Use sparse stored values/formulas, not the worksheet's styled rectangle.

    ``iter_rows()`` over ``max_row/max_column`` would allocate the inflated
    rectangle before applying our limits. The loaded worksheet's sparse cell
    map includes styled blanks and merged placeholders; neither defines a data
    boundary. Zero, False, whitespace strings, errors and formulas still do.
    A populated merged origin retains the full merged extent and evidence.
    """
    populated = {
        coordinates
        for coordinates, cell in sheet._cells.items()
        if cell.value is not None or cell.data_type == "f"
    }
    rows = max((row for row, _ in populated), default=0)
    columns = max((column for _, column in populated), default=0)
    for merged in sheet.merged_cells.ranges:
        if (merged.min_row, merged.min_col) in populated:
            rows = max(rows, merged.max_row)
            columns = max(columns, merged.max_col)
    return rows, columns


def _xlsx(body: bytes, result: _Document):
    import openpyxl
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        # No VBA preservation, formula evaluation or external workbook reads.
        workbook = openpyxl.load_workbook(io.BytesIO(body), data_only=False, keep_links=False)
        try:
            cached = openpyxl.load_workbook(
                io.BytesIO(body), read_only=True, data_only=True, keep_links=False
            )
            try:
                result.title(workbook.properties.title)
                if len(workbook.worksheets) > MAX_SHEETS:
                    result.issue(
                        "SHEET_LIMIT", "Spreadsheet sheets were truncated.", truncated=True
                    )
                for sheet in workbook.worksheets[:MAX_SHEETS]:
                    if (
                        sheet.sheet_state != "visible"
                        or any(dim.hidden for dim in sheet.row_dimensions.values())
                        or any(dim.hidden for dim in sheet.column_dimensions.values())
                    ):
                        result.issue(
                            "HIDDEN_CONTENT_INCLUDED",
                            "Hidden sheets/rows/columns are included, not silently discarded.",
                        )
                    rows, columns = _sheet_size(*_xlsx_data_bounds(sheet), result)
                    if not rows or not columns:
                        continue
                    merged = _merge_map(
                        (
                            (r.min_row - 1, r.max_row, r.min_col - 1, r.max_col)
                            for r in sheet.merged_cells.ranges
                        ),
                        rows,
                        columns,
                        result,
                    )
                    cache_rows = cached[sheet.title].iter_rows(
                        min_row=1, max_row=rows, min_col=1, max_col=columns
                    )
                    references = {}
                    matrix = []
                    for ri, cache_row in enumerate(cache_rows):
                        line = []
                        for ci in range(columns):
                            origin = merged.get((ri, ci), (ri, ci))
                            if origin not in references:
                                cell = sheet.cell(origin[0] + 1, origin[1] + 1)
                                value = cell.value
                                if cell.data_type == "f":
                                    # Merged origins occur at their top-left before aliases.
                                    cached_value = cache_row[ci].value
                                    if cached_value is None:
                                        result.issue(
                                            "FORMULA_NO_CACHED_VALUE",
                                            "A formula has no cached value; "
                                            "formula text is retained without evaluation.",
                                        )
                                        value = value if isinstance(value, str) else ""
                                    else:
                                        result.issue(
                                            "FORMULA_CACHED_VALUE",
                                            "Formula cached values may be stale; "
                                            "formulas were not evaluated.",
                                        )
                                        value = cached_value
                                if cell.data_type == "e":
                                    result.issue(
                                        "SPREADSHEET_CELL_ERROR",
                                        "An Excel error value was retained as literal text.",
                                    )
                                references[origin] = result.block(
                                    value,
                                    f"xlsx:sheet={json.dumps(sheet.title, ensure_ascii=False)};"
                                    f"cell={cell.coordinate}",
                                )
                            line.append(references[origin])
                        matrix.append(line)
                    result.table(
                        f"xlsx:sheet={json.dumps(sheet.title, ensure_ascii=False)}", matrix
                    )
            finally:
                cached.close()
        finally:
            workbook.close()
    except (InvalidFileException, KeyError, ValueError, OSError, zipfile.BadZipFile) as exc:
        raise _MaterialError(
            "INVALID_XLSX", "The workbook structure or a cell value is invalid."
        ) from exc


def _xls_formula_cells(book, sheet_index: int, body: bytes) -> set[tuple[int, int]]:
    """Read BIFF formula record coordinates only; never decompile/evaluate formulas.

    xlrd deliberately exposes only cached cell values. Its bounded input buffer
    is inspected for record headers to label those values accurately.
    """
    position = book._sh_abs_posn[sheet_index]
    data = book.mem
    if data is None:
        # BIFF2-4 are eagerly read by xlrd even when on_demand=True.
        if body.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
            from xlrd.compdoc import CompDoc

            container = CompDoc(body, logfile=io.StringIO())
            data = container.get_named_stream("Workbook") or container.get_named_stream("Book")
            if data is None:
                raise _MaterialError("INVALID_XLS", "Missing workbook stream.")
            position -= book.base
        else:
            data = body
    formulas = set()
    records = 0
    while position + 4 <= len(data):
        opcode, length = struct.unpack_from("<HH", data, position)
        start, end = position + 4, position + 4 + length
        if end > len(data):
            raise _MaterialError("INVALID_XLS", "Truncated BIFF record.")
        if opcode in {0x0006, 0x0206, 0x0406} and length >= 4:
            formulas.add(struct.unpack_from("<HH", data, start))
        if opcode == 0x000A:  # Sheet EOF.
            break
        position = end
        records += 1
        if records > MAX_NODES:
            raise _Limit("XLS BIFF record count exceeds the limit.")
    return formulas


def _xls(body: bytes, result: _Document):
    import xlrd
    from xlrd.compdoc import CompDocError

    try:
        with xlrd.open_workbook(
            file_contents=body, formatting_info=True, on_demand=True, logfile=io.StringIO()
        ) as workbook:
            if workbook.nsheets > MAX_SHEETS:
                result.issue("SHEET_LIMIT", "XLS sheets were truncated.", truncated=True)
            for si in range(min(workbook.nsheets, MAX_SHEETS)):
                formulas = _xls_formula_cells(workbook, si, body)
                sheet = workbook.sheet_by_index(si)
                if (
                    sheet.visibility
                    or any(row.hidden for row in sheet.rowinfo_map.values())
                    or any(col.hidden for col in sheet.colinfo_map.values())
                ):
                    result.issue(
                        "HIDDEN_CONTENT_INCLUDED", "Hidden XLS sheets/rows/columns are included."
                    )
                rows, columns = _sheet_size(sheet.nrows, sheet.ncols, result)
                if not rows or not columns:
                    continue
                merged = _merge_map(sheet.merged_cells, rows, columns, result)
                references = {}
                matrix = []
                for ri in range(rows):
                    line = []
                    for ci in range(columns):
                        origin = merged.get((ri, ci), (ri, ci))
                        if origin not in references:
                            cell = sheet.cell(*origin)
                            value = cell.value
                            if origin in formulas:
                                if value == "" or cell.ctype in {
                                    xlrd.XL_CELL_EMPTY,
                                    xlrd.XL_CELL_BLANK,
                                }:
                                    result.issue(
                                        "FORMULA_NO_CACHED_VALUE",
                                        "An XLS formula has an empty or unavailable cached value; "
                                        "no value was inferred.",
                                    )
                                else:
                                    result.issue(
                                        "FORMULA_CACHED_VALUE",
                                        "XLS formula results are cached values and may be stale; "
                                        "no evaluation occurred.",
                                    )
                            if cell.ctype == xlrd.XL_CELL_DATE:
                                try:
                                    value = xlrd.xldate_as_datetime(value, workbook.datemode)
                                except (ValueError, OverflowError) as exc:
                                    raise _MaterialError(
                                        "INVALID_XLS_DATE",
                                        "An XLS date is outside the supported range.",
                                    ) from exc
                            elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                                value = bool(value)
                            elif cell.ctype == xlrd.XL_CELL_ERROR:
                                value = xlrd.error_text_from_code.get(value, "#UNKNOWN!")
                                result.issue(
                                    "SPREADSHEET_CELL_ERROR", "An XLS error value was retained."
                                )
                            references[origin] = result.block(
                                value,
                                f"xls:sheet={json.dumps(sheet.name, ensure_ascii=False)};"
                                f"row={origin[0] + 1};column={origin[1] + 1}",
                            )
                        line.append(references[origin])
                    matrix.append(line)
                result.table(f"xls:sheet={json.dumps(sheet.name, ensure_ascii=False)}", matrix)
                if workbook.on_demand:
                    workbook.unload_sheet(si)
    except (xlrd.XLRDError, CompDocError, struct.error, UnicodeDecodeError) as exc:
        raise _MaterialError("INVALID_XLS", "The binary workbook is invalid or encrypted.") from exc


def _docx(body: bytes, result: _Document):
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from lxml.etree import XMLSyntaxError

    try:
        document = Document(io.BytesIO(body))
        result.title(document.core_properties.title)
        root = document.element
        if root.xpath(".//w:vanish"):
            result.issue("HIDDEN_CONTENT_INCLUDED", "Hidden Word text is included.")
        if root.xpath(
            ".//w:ins | .//w:del | .//w:txbxContent | .//w:altChunk | "
            ".//w:footnoteReference | .//w:endnoteReference"
        ):
            result.issue(
                "WORD_COMPLEX_CONTENT",
                "Tracked changes, textboxes, embedded content or notes require review; "
                "body paragraphs/tables only.",
            )
        if root.xpath(".//w:headerReference | .//w:footerReference"):
            result.issue(
                "WORD_HEADERS_FOOTERS_OMITTED",
                "Headers and footers are not part of body extraction.",
            )
        for element in document.element.body:
            kind = element.tag.rsplit("}", 1)[-1]
            if kind == "p":
                paragraph = Paragraph(element, document)
                if paragraph.text.strip():
                    result.block(
                        paragraph.text,
                        "docx:word/document.xml:" + element.getroottree().getpath(element),
                    )
            elif kind == "tbl":
                _word_table(Table(element, document), result)
    except (PackageNotFoundError, KeyError, ValueError, XMLSyntaxError) as exc:
        raise _MaterialError(
            "INVALID_DOCX", "The Word document package or table structure is invalid."
        ) from exc


def _word_table(table, result: _Document, depth: int = 0):
    if depth >= 20:
        raise _Limit("Nested Word table depth exceeds the limit.")
    element = table._tbl
    location = "docx:word/document.xml:" + element.getroottree().getpath(element)
    rows = table.rows
    count = len(rows)
    columns = len(table.columns)
    count, columns = _sheet_size(count, columns, result)
    if not count or not columns:
        return
    references = {}
    nested = {}
    matrix = []
    for ri in range(count):
        row = rows[ri]
        row_location = "docx:word/document.xml:" + row._tr.getroottree().getpath(row._tr)
        line = []
        for _ in range(min(row.grid_cols_before, columns)):
            result.issue(
                "SPARSE_TABLE",
                "Missing Word grid cells are empty slots referencing the source row.",
            )
            line.append(result.block("", row_location))
        for cell in row.cells[: max(0, columns - len(line))]:
            # python-docx resolves horizontal/vertical aliases to their source tc.
            path = cell._tc.getroottree().getpath(cell._tc)
            if path not in references:
                references[path] = result.block(cell.text, "docx:word/document.xml:" + path)
                for child in cell.tables:
                    nested[child._tbl.getroottree().getpath(child._tbl)] = child
            line.append(references[path])
        while len(line) < columns:
            result.issue(
                "SPARSE_TABLE",
                "Missing Word grid cells are empty slots referencing the source row.",
            )
            line.append(result.block("", row_location))
        matrix.append(line)
    result.table(location, matrix)
    for child in nested.values():
        _word_table(child, result, depth + 1)


def _doc(body: bytes, result: _Document):
    """Legacy Word text extraction; antiword itself runs inside the host sandbox."""
    try:
        with tempfile.TemporaryDirectory(prefix="youju-doc-") as directory:
            path = os.path.join(directory, "material.doc")
            with open(path, "wb") as stream:
                stream.write(body)
            # File output is bounded by the host RLIMIT_FSIZE, never unbounded PIPE.
            with tempfile.TemporaryFile() as stdout, tempfile.TemporaryFile() as stderr:
                process = subprocess.Popen(
                    ["/usr/bin/antiword", "-t", "-m", "UTF-8.txt", "-w", "0", path],
                    stdin=subprocess.DEVNULL,
                    stdout=stdout,
                    stderr=stderr,
                    env={"PATH": "/usr/bin:/bin", "LANG": "C.UTF-8", "HOME": directory},
                    cwd=directory,
                    start_new_session=True,
                )
                try:
                    process.wait(timeout=ANTIWORD_TIMEOUT)
                except subprocess.TimeoutExpired:
                    # /usr/bin/antiword can be a shell wrapper: kill its converter
                    # child too, not merely the wrapper process.
                    with contextlib.suppress(ProcessLookupError):
                        os.killpg(process.pid, signal.SIGKILL)
                    process.wait()
                    raise
                if process.returncode != 0:
                    raise _MaterialError(
                        "DOC_CONVERSION_FAILED",
                        "antiword could not decode this legacy Word document.",
                    )
                stdout.seek(0)
                converted = stdout.read(MAX_DECOMPRESSED_BYTES + 1)
                if len(converted) > MAX_DECOMPRESSED_BYTES:
                    raise _Limit("Legacy Word converted output exceeds the limit.")
                try:
                    text = converted.decode("utf-8")
                except UnicodeDecodeError as exc:
                    raise _MaterialError(
                        "DOC_ENCODING_ERROR", "antiword output is not valid UTF-8."
                    ) from exc
                for index, line in enumerate(text.splitlines(), 1):
                    if line.strip():
                        result.block(line, f"doc:antiword-text-line={index}")
                result.issue(
                    "DOC_TEXT_ONLY",
                    "Legacy Word is converted to text lines; "
                    "original table cells and pagination are not inferred.",
                )
                if not result.data["blocks"]:
                    result.issue("DOC_NO_TEXT", "The Word converter produced no extractable text.")
    except subprocess.TimeoutExpired as exc:
        raise _MaterialError(
            "DOC_CONVERSION_TIMEOUT", "antiword exceeded its execution deadline.", truncated=True
        ) from exc
    except FileNotFoundError as exc:
        raise _MaterialError(
            "DOC_CONVERTER_UNAVAILABLE", "The sandbox does not provide /usr/bin/antiword."
        ) from exc
    except OSError as exc:
        raise _MaterialError(
            "DOC_CONVERSION_IO_ERROR",
            "The bounded Word conversion could not read or write its temporary files.",
        ) from exc
